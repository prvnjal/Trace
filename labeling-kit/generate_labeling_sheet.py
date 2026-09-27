#!/usr/bin/env python3
"""
Generate the TRACE labeling workbook for the ML team.

Reads priority thermal events from PostGIS and writes one XLSX with batch
sheets (default 5 x 30 events), prefilled evidence columns, an EO Browser
satellite-photo link per event, and a dropdown label column.

Run from the project root (where docker-compose.yml and .env live):

    pip install openpyxl psycopg2-binary
    python labeling-kit/generate_labeling_sheet.py --events 150 --per-batch 30

DB credentials come from .env (POSTGRES_DB / POSTGRES_USER /
POSTGRES_PASSWORD / POSTGRES_PORT) or the environment. Defaults:
thermal_intel / thermal / 5432.
"""
import argparse
import os
import sys
from datetime import timedelta

LABELS = [
    "ROUTINE_FLARE",
    "PERSISTENT_THERMAL_SOURCE",
    "CANDIDATE_INDUSTRIAL_FIRE",
    "LIKELY_WILDFIRE",
    "LIKELY_AG_BURNING",
    "NEEDS_REVIEW",
]

GUIDE_ROWS = [
    ("ROUTINE_FLARE",
     "Routine gas flaring at refineries / petrochemical plants.",
     "Same industrial site, night after night, steady moderate heat."),
    ("PERSISTENT_THERMAL_SOURCE",
     "Always-hot normal operations (steel, power, cement, brick kilns).",
     "Many days, steady heat, day AND night, at a plant. Not an emergency."),
    ("CANDIDATE_INDUSTRIAL_FIRE",
     "Looks like a real unintended fire at or near industry.",
     "Sudden start, short duration, big heat spike, unusual for the site."),
    ("LIKELY_WILDFIRE",
     "Forest / scrub fire.",
     "Forest or scrub land, far from facilities, multiple days, many detections."),
    ("LIKELY_AG_BURNING",
     "Field / stubble burning.",
     "Farmland, 1-3 days, mostly daytime, often Oct-Nov in NW India."),
    ("NEEDS_REVIEW",
     "Not sure. A valid, useful answer.",
     "Use this rather than guessing. A guess poisons the data; this does not."),
]


def load_dotenv(path=".env"):
    if not os.path.exists(path):
        return
    with open(path, encoding="utf-8") as fh:
        for line in fh:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, val = line.split("=", 1)
            val = val.strip().strip('"').strip("'")
            os.environ.setdefault(key.strip(), val)


def eo_browser_link(lat, lon, first_date):
    frm = (first_date - timedelta(days=1)).isoformat()
    to = (first_date + timedelta(days=1)).isoformat()
    return (
        "https://apps.sentinel-hub.com/eo-browser/?zoom=12"
        f"&lat={lat:.4f}&lng={lon:.4f}&themeId=DEFAULT-THEME&datasetId=S2L2A"
        f"&fromTime={frm}T00%3A00%3A00.000Z&toTime={to}T23%3A59%3A59.999Z"
    )


QUERY = """
SELECT e.event_code, e.centroid_lat, e.centroid_lon,
       e.first_detected::date AS first_d, e.last_detected::date AS last_d,
       e.duration_hours, e.detection_count,
       e.max_frp, e.mean_frp, e.satellites,
       f.name  AS fac_name, f.facility_type AS fac_type,
       e.facility_distance_m, e.facilities_within_1km,
       e.landuse_class, e.landuse_tag,
       COALESCE(d.day_dets, 0)   AS day_dets,
       COALESCE(d.night_dets, 0) AS night_dets
FROM thermal_events e
LEFT JOIN industrial_facilities f ON f.id = e.nearest_facility_id
LEFT JOIN (
    SELECT event_id,
           COUNT(*) FILTER (WHERE upper(daynight) = 'D') AS day_dets,
           COUNT(*) FILTER (WHERE upper(daynight) = 'N') AS night_dets
    FROM thermal_detections
    GROUP BY event_id
) d ON d.event_id = e.id
{not_labeled}
ORDER BY e.detection_count DESC, e.duration_hours DESC,
         e.max_frp DESC NULLS LAST
LIMIT %s
"""


def fetch_events(conn, limit):
    with conn.cursor() as cur:
        # Skip events that already have a human label, if the table exists.
        cur.execute("SELECT to_regclass('public.event_labels')")
        has_labels = cur.fetchone()[0] is not None
        extra_join = ""
        extra_where = ""
        if has_labels:
            extra_join = ("LEFT JOIN event_labels el "
                          "ON el.event_code = e.event_code")
            extra_where = "WHERE el.event_code IS NULL "
        query = QUERY.format(not_labeled=extra_join)
        if extra_where:
            query = query.replace("ORDER BY", extra_where + "ORDER BY", 1)
        cur.execute(query, (limit,))
        cols = [d[0] for d in cur.description]
        return [dict(zip(cols, row)) for row in cur.fetchall()]


def build_workbook(events, per_batch):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    header = [
        "event_code", "lat", "lon", "first_detected", "last_detected",
        "days", "detections", "satellites", "max_FRP_MW", "mean_FRP_MW",
        "day_dets", "night_dets", "night_pct",
        "nearest_facility", "facility_type", "facility_dist_km",
        "facilities_1km", "land_use", "satellite_photo",
        "label", "notes",
    ]
    widths = [16, 10, 10, 13, 13, 7, 10, 22, 11, 11, 9, 10, 9,
              28, 14, 14, 12, 14, 14, 26, 30]
    label_col = get_column_letter(header.index("label") + 1)

    hdr_fill = PatternFill("solid", fgColor="1F2937")
    hdr_font = Font(color="FFFFFF", bold=True, size=10)

    dv = DataValidation(
        type="list",
        formula1='"' + ",".join(LABELS) + '"',
        allow_blank=True,
        showDropDown=False,
        promptTitle="Pick a label",
        prompt="Choose the category that best fits the evidence. When unsure, use NEEDS_REVIEW.",
    )

    for bi in range(0, len(events), per_batch):
        chunk = events[bi:bi + per_batch]
        ws = wb.active if bi == 0 else wb.create_sheet()
        ws.title = f"Batch {bi // per_batch + 1}"
        ws.append(header)
        for cell in ws[1]:
            cell.fill = hdr_fill
            cell.font = hdr_font
            cell.alignment = Alignment(horizontal="center", vertical="center",
                                        wrap_text=True)
        for r, ev in enumerate(chunk, start=2):
            total_dn = (ev["day_dets"] or 0) + (ev["night_dets"] or 0)
            night_pct = (ev["night_dets"] / total_dn) if total_dn else None
            sats = ", ".join(ev["satellites"]) if ev["satellites"] else ""
            fac_km = (ev["facility_distance_m"] / 1000
                      if ev["facility_distance_m"] is not None else None)
            row = [
                ev["event_code"], ev["centroid_lat"], ev["centroid_lon"],
                ev["first_d"], ev["last_d"],
                round((ev["duration_hours"] or 0) / 24, 1),
                ev["detection_count"], sats,
                ev["max_frp"], ev["mean_frp"],
                ev["day_dets"], ev["night_dets"], night_pct,
                ev["fac_name"], ev["fac_type"], fac_km,
                ev["facilities_within_1km"],
                ev["landuse_class"],
                None,   # satellite_photo hyperlink set below
                None,   # label (teammate fills)
                None,   # notes
            ]
            ws.append(row)
            photo_cell = ws.cell(row=r, column=header.index("satellite_photo") + 1)
            photo_cell.hyperlink = eo_browser_link(
                ev["centroid_lat"], ev["centroid_lon"], ev["first_d"])
            photo_cell.value = "open photo"
            photo_cell.style = "Hyperlink"
            ws.cell(row=r, column=4).number_format = "yyyy-mm-dd"
            ws.cell(row=r, column=5).number_format = "yyyy-mm-dd"
            ws.cell(row=r, column=13).number_format = "0%"
            for ci in (9, 10):
                ws.cell(row=r, column=ci).number_format = "0.00"
            ws.cell(row=r, column=16).number_format = "0.0"

        last_row = len(chunk) + 1
        ws.add_data_validation(dv)
        dv.add(f"{label_col}2:{label_col}{last_row}")
        for idx, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_properties.pageSetUpPr.fitToPage = True

    # Quick-reference guide sheet
    guide = wb.create_sheet("Label guide")
    guide.append(["label", "what it is", "pick when you see"])
    for cell in guide[1]:
        cell.fill = hdr_fill
        cell.font = hdr_font
    for lab, what, when in GUIDE_ROWS:
        guide.append([lab, what, when])
    guide.column_dimensions["A"].width = 28
    guide.column_dimensions["B"].width = 52
    guide.column_dimensions["C"].width = 62
    for row in guide.iter_rows(min_row=2):
        row[0].font = Font(bold=True)
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    guide.freeze_panes = "A2"
    return wb


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--events", type=int, default=150,
                    help="how many priority events to include (default 150)")
    ap.add_argument("--per-batch", type=int, default=30,
                    help="events per batch sheet (default 30)")
    ap.add_argument("--out", default="TRACE_labeling.xlsx",
                    help="output file (default TRACE_labeling.xlsx)")
    args = ap.parse_args()

    try:
        import psycopg2  # noqa: F401
    except ImportError:
        sys.exit("ERROR: psycopg2 not installed. Run: pip install psycopg2-binary")
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        sys.exit("ERROR: openpyxl not installed. Run: pip install openpyxl")

    import psycopg2

    load_dotenv(".env")
    conn = psycopg2.connect(
        host=os.environ.get("POSTGRES_HOST", "localhost"),
        port=int(os.environ.get("POSTGRES_PORT", "5432")),
        dbname=os.environ.get("POSTGRES_DB", "thermal_intel"),
        user=os.environ.get("POSTGRES_USER", "thermal"),
        password=os.environ.get("POSTGRES_PASSWORD", ""),
    )
    try:
        events = fetch_events(conn, args.events)
    finally:
        conn.close()

    if not events:
        sys.exit("No unlabeled events found. Nothing to write.")

    wb = build_workbook(events, args.per_batch)
    wb.save(args.out)
    print(f"Wrote {args.out}: {len(events)} events in "
          f"{(len(events) + args.per_batch - 1) // args.per_batch} batch sheets.")


if __name__ == "__main__":
    main()
