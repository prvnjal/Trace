#!/usr/bin/env python3
"""
Import the team's filled labeling workbook back into PostGIS.

Reads every "Batch *" sheet, takes each row with a non-blank label, and
upserts it into the event_labels table (one row per event per labeler).

Usage:
    python labeling-kit/import_labels.py TRACE_labeling_filled.xlsx ^
        --by "Batch 1=Aarav;Batch 2=Diya;Batch 3=Ishaan;Batch 4=Meera;Batch 5=Kabir"

If --by is omitted, the sheet name is used as the labeler name.

DB credentials come from .env, same as generate_labeling_sheet.py.
"""
import argparse
import os
import sys

LABELS = {
    "ROUTINE_FLARE",
    "PERSISTENT_THERMAL_SOURCE",
    "CANDIDATE_INDUSTRIAL_FIRE",
    "LIKELY_WILDFIRE",
    "LIKELY_AG_BURNING",
    "NEEDS_REVIEW",
}

UPSERT = """
INSERT INTO event_labels (event_code, label, labeled_by, notes)
VALUES (%s, %s, %s, %s)
ON CONFLICT (event_code, labeled_by)
DO UPDATE SET label = EXCLUDED.label,
              notes = EXCLUDED.notes,
              labeled_at = CURRENT_TIMESTAMP
"""


def load_dotenv(path=".env"):
    if not os.path.exists(path):
        return
    with open(path, encoding="utf-8") as fh:
        for line in fh:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, val = line.split("=", 1)
            val = val.strip().strip('"').strip("'")
            os.environ.setdefault(key.strip(), val)


def parse_by(spec):
    mapping = {}
    if spec:
        for part in spec.split(";"):
            if "=" in part:
                sheet, name = part.split("=", 1)
                mapping[sheet.strip()] = name.strip()
    return mapping


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", help="filled labeling workbook")
    ap.add_argument("--by", default="",
                    help='sheet->name map, e.g. "Batch 1=Aarav;Batch 2=Diya"')
    args = ap.parse_args()

    try:
        import openpyxl  # noqa: F401
    except ImportError:
        sys.exit("ERROR: openpyxl not installed. Run: pip install openpyxl")
    try:
        import psycopg2  # noqa: F401
    except ImportError:
        sys.exit("ERROR: psycopg2 not installed. Run: pip install psycopg2-binary")

    import openpyxl
    import psycopg2

    by = parse_by(args.by)
    wb = openpyxl.load_workbook(args.xlsx, data_only=False)

    rows, errors = [], []
    for ws in wb.worksheets:
        if not ws.title.startswith("Batch"):
            continue
        header = [c.value for c in ws[1]]
        try:
            i_code = header.index("event_code")
            i_label = header.index("label")
            i_notes = header.index("notes")
        except ValueError:
            errors.append(f"{ws.title}: missing event_code/label/notes columns")
            continue
        who = by.get(ws.title, ws.title)
        for r in ws.iter_rows(min_row=2, values_only=True):
            code = r[i_code]
            label = (r[i_label] or "").strip().upper() if r[i_label] else ""
            if not code or not label:
                continue
            if label not in LABELS:
                errors.append(f"{ws.title} {code}: bad label '{r[i_label]}'")
                continue
            notes = r[i_notes] if r[i_notes] else None
            rows.append((code, label, who, notes))

    if errors:
        print("ERRORS — fix these in the workbook and re-run:")
        for e in errors:
            print("  -", e)
        sys.exit(1)
    if not rows:
        sys.exit("No labeled rows found. Nothing imported.")

    load_dotenv(".env")
    conn = psycopg2.connect(
        host=os.environ.get("POSTGRES_HOST", "localhost"),
        port=int(os.environ.get("POSTGRES_PORT", "5432")),
        dbname=os.environ.get("POSTGRES_DB", "thermal_intel"),
        user=os.environ.get("POSTGRES_USER", "thermal"),
        password=os.environ.get("POSTGRES_PASSWORD", ""),
    )
    try:
        with conn.cursor() as cur:
            cur.execute("SELECT to_regclass('public.event_labels')")
            if cur.fetchone()[0] is None:
                sys.exit("ERROR: event_labels table missing. "
                         "Apply 004_event_labels.sql first.")
            cur.execute("SELECT event_code FROM thermal_events")
            existing = {r[0] for r in cur.fetchall()}
            skipped = [r for r in rows if r[0] not in existing]
            rows = [r for r in rows if r[0] in existing]
            if skipped:
                print(f"WARNING: {len(skipped)} labels skipped - "
                      "events aged out of the refresh window:")
                for _code, _lab, _who, _ in skipped:
                    print(f"  - {_code} ({_who}: {_lab})")
            cur.executemany(UPSERT, rows)
        conn.commit()
    finally:
        conn.close()

    from collections import Counter
    dist = Counter(label for _, label, _, _ in rows)
    print(f"Imported {len(rows)} labels.")
    for lab in sorted(dist):
        print(f"  {lab:28s} {dist[lab]}")


if __name__ == "__main__":
    main()
