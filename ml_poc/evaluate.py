"""Evaluate the trained model against the human-labeled workbook rows.

The workbook is a SELF-CONTAINED evaluation set: features and labels come
from the file itself, so this score is immune to the refresh pipeline
reshuffling event codes in PostGIS.

- NEEDS_REVIEW rows are excluded from scoring (counted separately).
- Reports accuracy, macro/weighted F1, per-class precision/recall/F1,
  confusion matrix, and a majority-class baseline for comparison.
- Also reports weak-label vs human agreement on events carrying both
  (a diagnostic of heuristic quality, not of the model).

Usage:  python ml_poc/evaluate.py [workbook.xlsx]
Writes ml_poc/outputs/eval_report.json and scorecard.txt.

Run from the project root.
"""
import csv
import json
import os
import pickle
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from db import LABELS5, MODEL_VERSION, NUMERIC_COLS, derive_features

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.join(HERE, "outputs")
MODEL_PKL = os.path.join(OUT_DIR, "model.pkl")
WEAK_CSV = os.path.join(OUT_DIR, "ml_weak_labels.csv")
REPORT_JSON = os.path.join(OUT_DIR, "eval_report.json")
SCORECARD_TXT = os.path.join(OUT_DIR, "scorecard.txt")

LABEL2IDX = {lab: i for i, lab in enumerate(LABELS5)}
IDX2LABEL = {i: lab for lab, i in LABEL2IDX.items()}


def parse_n_sats(v):
    if v is None or str(v).strip() == "":
        return 0
    return len([s for s in str(v).split(",") if s.strip()])


def fnum(v):
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def load_eval_rows(xlsx_path):
    from openpyxl import load_workbook
    wb = load_workbook(xlsx_path, data_only=True)
    rows, n_needreview, n_bad = [], 0, 0
    for ws in wb.worksheets:
        if not ws.title.startswith("Batch"):
            continue
        header = [c.value for c in ws[1]]
        for r in ws.iter_rows(min_row=2, values_only=True):
            rec = {header[i]: r[i] for i in range(len(header))}
            label = (rec.get("label") or "").strip()
            if label == "NEEDS_REVIEW":
                n_needreview += 1
                continue
            if label not in LABEL2IDX:
                n_bad += 1
                continue
            canon = {
                "days": fnum(rec.get("days")),
                "detections": fnum(rec.get("detections")),
                "n_satellites": parse_n_sats(rec.get("satellites")),
                "max_frp": fnum(rec.get("max_FRP_MW")),
                "mean_frp": fnum(rec.get("mean_FRP_MW")),
                "day_dets": fnum(rec.get("day_dets")),
                "night_dets": fnum(rec.get("night_dets")),
                "facility_dist_km": fnum(rec.get("facility_dist_km")),
                "facilities_1km": fnum(rec.get("facilities_1km")),
                "land_use": rec.get("land_use"),
            }
            feat = derive_features(canon)
            feat["event_code"] = rec.get("event_code")
            feat["human_label"] = label
            rows.append(feat)
    return rows, n_needreview, n_bad


def build_matrix(rows, feature_cols, landuse_cats):
    import pandas as pd
    df = pd.DataFrame(rows)
    X_num = df[NUMERIC_COLS].apply(pd.to_numeric, errors="coerce")
    dummies = pd.get_dummies(df["land_use"].fillna("unknown"), prefix="lu")
    dummies = dummies.reindex(columns=landuse_cats, fill_value=0)
    X = pd.concat([X_num, dummies], axis=1)
    return X.reindex(columns=feature_cols, fill_value=0)


def main():
    xlsx = sys.argv[1] if len(sys.argv) > 1 else "TRACE_labeling_filled.xlsx"
    if not os.path.exists(xlsx):
        sys.exit(f"ERROR: workbook not found: {xlsx}\n"
                 "Copy TRACE_labeling_filled.xlsx into the project root first.")
    with open(MODEL_PKL, "rb") as fh:
        bundle = pickle.load(fh)
    model = bundle["model"]

    eval_rows, n_needreview, n_bad = load_eval_rows(xlsx)
    X = build_matrix(eval_rows, bundle["feature_cols"], bundle["landuse_cats"])
    y_true = [LABEL2IDX[r["human_label"]] for r in eval_rows]
    raw_pred = list(model.predict(X))
    # Map through the model's own class list: if a class had no training
    # data it is absent from bundle["classes"], and raw indices would
    # otherwise misalign with the 5-class LABEL2IDX.
    y_pred = [LABEL2IDX[bundle["classes"][i]] for i in raw_pred]

    from sklearn.metrics import (accuracy_score, f1_score,
                                 classification_report, confusion_matrix)
    from collections import Counter
    acc = accuracy_score(y_true, y_pred)
    macro_f1 = f1_score(y_true, y_pred, average="macro", zero_division=0)
    weighted_f1 = f1_score(y_true, y_pred, average="weighted", zero_division=0)
    report = classification_report(y_true, y_pred, labels=list(range(len(LABELS5))),
                                   target_names=LABELS5, output_dict=True,
                                   zero_division=0)
    cm = confusion_matrix(y_true, y_pred, labels=list(range(len(LABELS5)))).tolist()

    # Majority-class baseline (most common TRAIN label -> predict everywhere).
    with open(os.path.join(OUT_DIR, "train_report.json"), encoding="utf-8") as fh:
        train_rep = json.load(fh)
    maj_lab = max(train_rep["train_label_dist"],
                  key=train_rep["train_label_dist"].get)
    maj_idx = LABEL2IDX[maj_lab]
    baseline_acc = sum(1 for t in y_true if t == maj_idx) / len(y_true)

    # Diagnostic: weak-label vs human agreement where both exist.
    weak = {}
    if os.path.exists(WEAK_CSV):
        with open(WEAK_CSV, newline="", encoding="utf-8") as fh:
            for r in csv.DictReader(fh):
                weak[r["event_code"]] = r["weak_label"]
    agree_n = agree_ok = 0
    for r in eval_rows:
        w = weak.get(r["event_code"])
        if w:
            agree_n += 1
            agree_ok += int(w == r["human_label"])

    eval_report = {
        "model_version": bundle["version"],
        "flavor": bundle["flavor"],
        "n_eval": len(eval_rows),
        "n_needreview_excluded": n_needreview,
        "n_invalid_label_rows": n_bad,
        "accuracy": round(acc, 4),
        "macro_f1": round(macro_f1, 4),
        "weighted_f1": round(weighted_f1, 4),
        "baseline_majority_label": maj_lab,
        "baseline_accuracy": round(baseline_acc, 4),
        "per_class": {
            lab: {k: round(v, 4) for k, v in report[lab].items()
                  if k in ("precision", "recall", "f1-score", "support")}
            for lab in LABELS5
        },
        "confusion_matrix": cm,
        "confusion_matrix_labels": LABELS5,
        "weak_vs_human_agreement": {
            "n_compared": agree_n,
            "agreement": round(agree_ok / agree_n, 4) if agree_n else None,
        },
        "eval_label_dist": {lab: sum(1 for r in eval_rows
                                     if r["human_label"] == lab)
                            for lab in LABELS5},
    }
    with open(REPORT_JSON, "w", encoding="utf-8") as fh:
        json.dump(eval_report, fh, indent=2)

    # ---- human-readable scorecard ----
    L = []
    A = L.append
    A("=" * 64)
    A(f"TRACE ML PoC scorecard  (model {bundle['version']}, {bundle['flavor']})")
    A("=" * 64)
    A(f"Eval set : {len(eval_rows)} human-labeled events from the labeling workbook")
    A(f"           ({n_needreview} NEEDS_REVIEW excluded from scoring)")
    A(f"Train set: {train_rep['n_train']} weak-labeled current events "
      f"(heuristics, not ground truth)")
    A("")
    A(f"Accuracy          : {acc:.3f}   (baseline always-{maj_lab}: {baseline_acc:.3f})")
    A(f"Macro F1          : {macro_f1:.3f}")
    A(f"Weighted F1       : {weighted_f1:.3f}")
    A("")
    A(f"{'class':28s} {'prec':>6s} {'rec':>6s} {'f1':>6s} {'n':>4s}")
    A("-" * 64)
    for lab in LABELS5:
        c = eval_report["per_class"][lab]
        A(f"{lab:28s} {c['precision']:6.3f} {c['recall']:6.3f} "
          f"{c['f1-score']:6.3f} {int(c['support']):4d}")
    A("")
    A("Confusion matrix (rows = human label, cols = predicted):")
    hdr = " " * 16 + "".join(f"{l[:10]:>11s}" for l in LABELS5)
    A(hdr)
    for i, lab in enumerate(LABELS5):
        A(f"{lab[:15]:15s} " + "".join(f"{cm[i][j]:11d}" for j in range(len(LABELS5))))
    A("")
    if agree_n:
        A(f"Weak-label vs human agreement (on {agree_n} events carrying both): "
          f"{agree_ok / agree_n:.3f}")
        A("  (heuristic quality diagnostic — not a model score)")
    A("")
    A("Caveats: weak labels are noisy by design; ROUTINE_FLARE has only "
      f"{eval_report['eval_label_dist']['ROUTINE_FLARE']} eval examples so its "
      "recall estimate is noisy; event codes churn across refreshes, so this "
      "score is a point-in-time evaluation on the frozen workbook, not on the "
      "live event table.")
    scorecard = "\n".join(L)
    with open(SCORECARD_TXT, "w", encoding="utf-8") as fh:
        fh.write(scorecard + "\n")
    print(scorecard)
    print(f"\nSaved {REPORT_JSON} and {SCORECARD_TXT}")


if __name__ == "__main__":
    main()
