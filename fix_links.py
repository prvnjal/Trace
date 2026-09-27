"""Fix v2: correct Copernicus dataset ID (S2_L2A_CDAS) in photo links."""
import openpyxl
from datetime import timedelta

wb = openpyxl.load_workbook("TRACE_labeling.xlsx")
fixed = 0
for ws in wb.worksheets:
    if not ws.title.startswith("Batch"):
        continue
    header = [c.value for c in ws[1]]
    col = header.index("satellite_photo") + 1
    lat_c = header.index("lat") + 1
    lon_c = header.index("lon") + 1
    date_c = header.index("first_detected") + 1
    for r in range(2, ws.max_row + 1):
        lat = ws.cell(row=r, column=lat_c).value
        lon = ws.cell(row=r, column=lon_c).value
        d = ws.cell(row=r, column=date_c).value
        if lat is None or d is None:
            continue
        frm = (d - timedelta(days=1)).strftime("%Y-%m-%d")
        to = (d + timedelta(days=1)).strftime("%Y-%m-%d")
        url = (
            "https://browser.dataspace.copernicus.eu/?zoom=12"
            f"&lat={lat:.4f}&lng={lon:.4f}&themeId=DEFAULT-THEME"
            f"&datasetId=S2_L2A_CDAS&layerId=1_TRUE_COLOR&dateMode=SINGLE"
            f"&fromTime={frm}T00%3A00%3A00.000Z&toTime={to}T23%3A59%3A59.999Z"
        )
        cell = ws.cell(row=r, column=col)
        cell.hyperlink = url
        cell.value = "open photo"
        cell.style = "Hyperlink"
        fixed += 1
wb.save("TRACE_labeling.xlsx")
print(f"Fixed {fixed} links with S2_L2A_CDAS.")
